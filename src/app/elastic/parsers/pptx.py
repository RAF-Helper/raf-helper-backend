"""PPTX parser using python-pptx."""

import io
from pathlib import Path
from typing import Any

from pptx import Presentation

from .base import BaseParser


class PptxParser(BaseParser):
    """Parser for PowerPoint (.pptx) presentations."""

    def parse_file(self, file_path: str | Path) -> dict[str, Any]:
        with open(file_path, "rb") as f:
            return self._parse(f.read(), str(Path(file_path).name))

    def parse_bytes(self, data: bytes, filename: str | None = None) -> dict[str, Any]:
        return self._parse(data, filename or "presentation.pptx")

    def _parse(self, data: bytes, filename: str) -> dict[str, Any]:
        prs = Presentation(io.BytesIO(data))
        slides_text: list[str] = []

        for slide in prs.slides:
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                slides_text.append("\n".join(texts))

        full_text = "\n\n".join(slides_text)
        return {
            "text": full_text,
            "metadata": {
                "filename": filename,
                "slide_count": len(prs.slides),
            },
        }
