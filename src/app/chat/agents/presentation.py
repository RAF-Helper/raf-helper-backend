"""Presentation agent — generates PowerPoint slides from context materials.

Sub-agents (all stubs):
  Planner → Slider → SlideSummarizer (iterative per slide)
Output: real .pptx file via python-pptx.
"""

import io
import uuid
from pathlib import Path
from typing import Optional

from pptx import Presentation as PPTXPresentation

from ..state import AgentInput, AgentOutput
from ...schemas.agent import AgentType
from ...repositories.elasticsearch_repository import ElasticsearchRepository
from ...config import settings
from .base import BaseAgent

_PLANNER_PROMPT = """Based on the materials and user's topic, create a slide plan.
For each slide: title + 3-5 bullet points. Return as numbered list."""

_SLIDER_PROMPT = """Write content for the next slide using context materials and
summary of previous slides. Return:\n  Title: <title>\n  Bullets:\n  - ..."""

_SLIDE_SUMMARIZER_PROMPT = "Summarize the slides written so far into a brief paragraph for context in next slides."

ROUTING_EXAMPLES = [
    "сделай презентацию", "создай слайды", "сделай слайд-шоу",
    "presentation", "make slides", "generate powerpoint",
    "презентация по теме", "PowerPoint из материалов",
]


class PresentationAgent(BaseAgent):
    agent_type = AgentType.presentation

    def __init__(self, es_repo: Optional[ElasticsearchRepository] = None):
        self.es_repo = es_repo or ElasticsearchRepository()

    async def run(self, inp: AgentInput) -> AgentOutput:
        hits = await self.es_repo.search_context(
            query=inp.masked_message, document_ids=inp.file_ids or None, limit=20
        )
        context_chunks = [h.get("chunk_text", "") for h in hits if h.get("chunk_text")]

        # TODO Stage 1 — Planner:
        # plan_text = await self._call_llm(_PLANNER_PROMPT, self._build_messages([], inp.masked_message, context_chunks))
        # slides_plan = _parse_slide_plan(plan_text)
        slides_plan = [
            {"title": "Introduction (stub)", "bullets": ["Point 1", "Point 2"]},
            {"title": "Main Topic (stub)", "bullets": ["Point A", "Point B", "Point C"]},
            {"title": "Conclusion (stub)", "bullets": ["Summary", "Next steps"]},
        ]

        # TODO Stage 2 — Slider (iterative with slide summaries for context)
        # TODO Stage 3 — SlideSummarizer after each slide

        pptx_path = self._build_pptx(slides_plan, inp.chat_id)

        content = (
            f"[STUB] PresentationAgent: LLM integration pending.\n"
            f"Would generate {len(slides_plan)} slides. Skeleton saved to: {pptx_path}"
        )

        return AgentOutput(content=content, agent_used=self.agent_type, masked_query=inp.masked_message,
                           metadata={"slide_count": len(slides_plan), "pptx_path": str(pptx_path)})

    def _build_pptx(self, slides: list[dict], chat_id: int) -> Path:
        prs = PPTXPresentation()
        layout = prs.slide_layouts[1]
        for slide_data in slides:
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = slide_data["title"]
            tf = slide.placeholders[1].text_frame
            for i, bullet in enumerate(slide_data.get("bullets", [])):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = bullet
        out_dir = Path(settings.storage_path) / "presentations"
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / f"chat_{chat_id}_{uuid.uuid4().hex[:8]}.pptx"
        buf = io.BytesIO()
        prs.save(buf)
        out_path.write_bytes(buf.getvalue())
        return out_path
