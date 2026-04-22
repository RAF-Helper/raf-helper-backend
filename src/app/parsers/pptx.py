"""PPTX parser using python-pptx."""

from pathlib import Path
from uuid import uuid4

from pptx import Presentation

from .base import BaseParser
from ..schemas.document import Document, DocumentChunk, DocumentFormat

class PptxParser(BaseParser[Presentation]):
    """Parser for PowerPoint (.pptx) presentations."""

    def load(self, filename: str) -> Presentation:
        """Load PPTX file into a Presentation object."""
        return Presentation(filename)

    def parse(self, artifact: Presentation, filename: str) -> Document:
        """Parse PPTX into a Document model, treating each slide as a block."""
        chunks: list[DocumentChunk] = []

        for i, slide in enumerate(artifact.slides):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())

            if texts:
                slide_text = "\n".join(texts)
                chunks.append(DocumentChunk(
                    content=slide_text,
                    level1=f"Slide {i+1}",
                    level2=None,
                    level3=None,
                    metadata={"page_number": i + 1}
                ))

        return Document(
            id=uuid4(),
            name=Path(filename).stem,
            filename=Path(filename).name,
            format=DocumentFormat.PPTX,
            document_chunks=chunks
        )
