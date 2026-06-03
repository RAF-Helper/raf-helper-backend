import os
import logging
from pptx import Presentation
from pptx.util import Inches, Pt

logger = logging.getLogger(__name__)

class FileExportService:
    """Service to handle exporting generated content to physical files."""

    def __init__(self, storage_path: str = "./storage/exports"):
        self.storage_path = storage_path
        os.makedirs(self.storage_path, exist_ok=True)

    def export_latex(self, filename: str, content: str) -> str:
        """Export content to a LaTeX .tex file."""
        filepath = os.path.join(self.storage_path, f"{filename}.tex")

        # Basic LaTeX template
        template = (
            "\\documentclass{article}\n"
            "\\usepackage[utf8]{inputenc}\n"
            "\\usepackage[russian]{babel}\n"
            "\\begin{document}\n"
            f"{content}\n"
            "\\end{document}"
        )

        with open(filepath, "w", encoding="utf-8") as f:
            f.write(template)

        logger.info(f"LaTeX file exported to {filepath}")
        return filepath

    def export_pptx(self, filename: str, slides_content: list[dict]) -> str:
        """
        Export slides to a .pptx file.
        Expects slides_content as a list of dicts: [{"title": "...", "content": "..."}]
        """
        prs = Presentation()

        for slide_data in slides_content:
            slide_layout = prs.slide_layouts[1] # Title and Content layout
            slide = prs.slides.add_slide(slide_layout)

            title = slide.shapes.title
            content = slide.placeholders[1]

            title.text = slide_data.get("title", "Untitled Slide")
            content.text = slide_data.get("content", "")

        filepath = os.path.join(self.storage_path, f"{filename}.pptx")
        prs.save(filepath)

        logger.info(f"PPTX file exported to {filepath}")
        return filepath

export_service = FileExportService()
